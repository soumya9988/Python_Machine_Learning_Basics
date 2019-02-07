import openpyxl
import random
wb = openpyxl.Workbook()
ws = wb.active
for i in range(1, 11):
    ws['A' + str(i)] = random.randint(5, 65)


ref_obj = openpyxl.chart.Reference(ws, min_col=1, min_row=1, max_col=1, max_row=10)
series_obj = openpyxl.chart.Series(ref_obj, title='Random Chart')
chart_obj = openpyxl.chart.LineChart()
chart_obj.title = 'Random Numbers'
chart_obj.append(series_obj)
ws.add_chart(chart_obj, 'E1')
wb.save('sample_Chart.xlsx')
